from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from django.core.management.base import BaseCommand

from tests_app.models import TestCategory, DiagnosticTest
from packages.models import HealthPackage


class Command(BaseCommand):
    help = "Imports tests, profiles and packages from tests.xlsx safely."

    def handle(self, *args, **options):

        # Project root where manage.py and tests.xlsx are located
        BASE_DIR = Path(__file__).resolve().parents[3]
        excel_path = BASE_DIR / "tests.xlsx"

        if not excel_path.exists():
            self.stdout.write(
                self.style.ERROR(
                    f"tests.xlsx not found here: {excel_path}"
                )
            )
            return

        self.stdout.write("Opening tests.xlsx...")

        workbook = load_workbook(
            excel_path,
            data_only=True,
            read_only=True
        )

        sheet_name = "PDF_All_Data"

        if sheet_name not in workbook.sheetnames:
            self.stdout.write(
                self.style.ERROR(
                    f"Sheet '{sheet_name}' not found."
                )
            )
            return

        sheet = workbook[sheet_name]

        # Read header row
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))

        header_map = {
            str(header).strip(): index
            for index, header in enumerate(headers)
            if header is not None
        }

        required_columns = [
            "Test name",
            "Deal type",
            "Deal type id",
            "MRP",
        ]

        missing_columns = [
            column
            for column in required_columns
            if column not in header_map
        ]

        if missing_columns:
            self.stdout.write(
                self.style.ERROR(
                    f"Missing Excel columns: {', '.join(missing_columns)}"
                )
            )
            return

        tests_created = 0
        tests_existing = 0

        packages_created = 0
        packages_existing = 0

        skipped_invalid = 0

        seen_test_codes = set()
        seen_package_codes = set()

        # Categories for Excel data
        parameter_category, _ = TestCategory.objects.get_or_create(
            name="Parameters",
            defaults={"slug": "parameters"}
        )

        profile_category, _ = TestCategory.objects.get_or_create(
            name="Profiles",
            defaults={"slug": "profiles"}
        )

        for row in sheet.iter_rows(min_row=2, values_only=True):

            name = row[header_map["Test name"]]
            deal_type = row[header_map["Deal type"]]
            deal_type_id = row[header_map["Deal type id"]]
            mrp = row[header_map["MRP"]]

            if not name or not deal_type or not deal_type_id:
                skipped_invalid += 1
                continue

            name = str(name).strip()
            deal_type = str(deal_type).strip().lower()
            deal_type_id = str(deal_type_id).strip()

            if not name or not deal_type_id:
                skipped_invalid += 1
                continue

            try:
                price = Decimal(str(mrp))
            except (InvalidOperation, TypeError, ValueError):
                skipped_invalid += 1
                continue

            # ==========================
            # PARAMETERS AND PROFILES
            # ==========================

            if deal_type in ("parameter", "profile"):

                test_code = f"XL-{deal_type_id}"

                if len(test_code) > 20:
                    self.stdout.write(
                        self.style.WARNING(
                            f"Skipped code longer than 20 characters: {test_code}"
                        )
                    )
                    skipped_invalid += 1
                    continue

                if test_code in seen_test_codes:
                    continue

                seen_test_codes.add(test_code)

                category = (
                    parameter_category
                    if deal_type == "parameter"
                    else profile_category
                )

                _, created = DiagnosticTest.objects.get_or_create(
                    test_code=test_code,
                    defaults={
                        "name": name[:255],
                        "category": category,
                        "price": price,
                        "is_available": True,
                    }
                )

                if created:
                    tests_created += 1
                else:
                    tests_existing += 1

            # ==========================
            # PACKAGES
            # ==========================

            elif deal_type == "package":

                package_code = f"PKG-{deal_type_id}"

                if len(package_code) > 20:
                    self.stdout.write(
                        self.style.WARNING(
                            f"Skipped package code longer than 20 characters: "
                            f"{package_code}"
                        )
                    )
                    skipped_invalid += 1
                    continue

                if package_code in seen_package_codes:
                    continue

                seen_package_codes.add(package_code)

                _, created = HealthPackage.objects.get_or_create(
                    package_code=package_code,
                    defaults={
                        "name": name[:255],
                        "price": price,
                        "is_available": True,
                    }
                )

                if created:
                    packages_created += 1
                else:
                    packages_existing += 1

        self.stdout.write("")
        self.stdout.write(
            self.style.SUCCESS("========== IMPORT COMPLETED ==========")
        )

        self.stdout.write(
            f"New tests/profiles: {tests_created}"
        )

        self.stdout.write(
            f"Existing tests/profiles skipped: {tests_existing}"
        )

        self.stdout.write(
            f"New packages: {packages_created}"
        )

        self.stdout.write(
            f"Existing packages skipped: {packages_existing}"
        )

        self.stdout.write(
            f"Invalid rows skipped: {skipped_invalid}"
        )

        self.stdout.write(
            f"Total DiagnosticTest records: "
            f"{DiagnosticTest.objects.count()}"
        )

        self.stdout.write(
            f"Total HealthPackage records: "
            f"{HealthPackage.objects.count()}"
        )